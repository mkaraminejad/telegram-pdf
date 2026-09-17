"""Unit Tests for DOCX OpenXML RTL and BiDi Attributes.

Ensures that generated Word documents contain genuine OpenXML XML elements:
1. `<w:bidi/>` on paragraphs
2. `<w:rtl/>` on runs
3. `<w:rFonts w:cs="..."/>` for Persian complex script fonts
4. `<w:bidiVisual/>` on tables
"""
import unittest
import tempfile
import os
from docx import Document
from bot.services.docx_builder import DocxBuilder


class TestDocxRTL(unittest.TestCase):

    def setUp(self):
        self.builder = DocxBuilder(default_font_fa="Vazirmatn", default_font_size=12)
        self.temp_dir = tempfile.mkdtemp()
        self.output_docx = os.path.join(self.temp_dir, "test_output.docx")

    def tearDown(self):
        if os.path.exists(self.output_docx):
            os.remove(self.output_docx)
        if os.path.exists(self.temp_dir):
            os.rmdir(self.temp_dir)

    def test_paragraph_bidi_xml(self):
        """Paragraph must contain <w:bidi/> tag in OpenXML."""
        doc = self.builder.create_document()
        p = self.builder.add_rtl_paragraph(doc, "این یک پاراگراف فارسی برای تست است.")

        xml_str = p._p.xml
        self.assertIn("w:bidi", xml_str, "Paragraph properties missing <w:bidi/> tag")
        self.assertIn('w:jc w:val="right"', xml_str, "Paragraph alignment must be right")

    def test_run_rtl_xml(self):
        """Persian runs must contain <w:rtl/> and complex script font attributes."""
        doc = self.builder.create_document()
        p = self.builder.add_rtl_paragraph(doc, "متن فارسی با کلمه English و اعداد ۱۲۳")

        xml_str = p._p.xml
        self.assertIn("w:rtl", xml_str, "Run properties missing <w:rtl/> tag")
        self.assertIn('w:cs="Vazirmatn"', xml_str, "Run missing Vazirmatn complex script font")

    def test_table_bidi_visual_xml(self):
        """Table must contain <w:bidiVisual/> in tblPr to render columns Right-to-Left."""
        doc = self.builder.create_document()
        table_data = [
            ["ردیف", "نام محصول", "قیمت (تومان)"],
            ["۱", "کتاب آموزش پایتون", "۲۵۰,۰۰۰"],
            ["۲", "دوره هوش مصنوعی", "۹۰۰,۰۰۰"]
        ]
        table = self.builder.add_rtl_table(doc, table_data)

        table_xml = table._tbl.xml
        self.assertIn("w:bidiVisual", table_xml, "Table missing <w:bidiVisual/> RTL property")

    def test_full_document_assembly(self):
        """Test compiling extracted content structures into a valid docx file."""
        content_pages = [
            {
                "page_number": 1,
                "is_scanned": False,
                "blocks": [
                    {"type": "heading", "level": 1, "text": "فصل اول: مقدمه"},
                    {"type": "paragraph", "text": "این اولین پاراگراف از کتاب فارسی است."},
                    {"type": "bullet", "text": "نکته اول درباره الگوریتم BiDi"},
                    {"type": "bullet", "text": "نکته دوم درباره Tesseract OCR"},
                    {
                        "type": "table",
                        "data": [
                            ["ستون ۱", "ستون ۲"],
                            ["داده الف", "داده ب"]
                        ]
                    }
                ]
            }
        ]

        saved_path = self.builder.build_from_extracted_content(content_pages, self.output_docx)
        self.assertTrue(os.path.exists(saved_path), "DOCX file was not saved")
        self.assertGreater(os.path.getsize(saved_path), 1000, "DOCX file size suspiciously small")

        # Reload with python-docx to verify format validity
        reopened_doc = Document(saved_path)
        self.assertGreater(len(reopened_doc.paragraphs), 0)
        self.assertEqual(len(reopened_doc.tables), 1)


if __name__ == "__main__":
    unittest.main()
